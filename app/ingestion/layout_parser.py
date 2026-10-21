"""Extract text from PDF, CSV, DOCX, and raster image files."""

import csv
from pathlib import Path
from typing import Any

from app.ingestion.ocr import extract_image_text


def ingest_document(file_path: str | Path, ocr_engine: Any = None) -> list[dict]:
    path = Path(file_path)
    extension = path.suffix.lower()
    supported_extensions = {".csv", ".docx", ".pdf", ".png", ".jpg", ".jpeg"}
    if extension not in supported_extensions:
        raise ValueError(f"Unsupported document type: {extension}")

    if extension == ".csv":
        return _ingest_csv(path)
    if extension == ".docx":
        return _ingest_docx(path)

    try:
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError("Document ingestion requires pillow.") from exc

    documents = []
    if extension == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("PDF ingestion requires pymupdf.") from exc
        with fitz.open(str(path)) as pdf:
            for page_number, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()
                if len(text) >= 30:
                    documents.append({"page": page_number, "text": text, "method": "pdf_text"})
                    continue
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                lines = extract_image_text(image, ocr_engine)
                documents.append({
                    "page": page_number,
                    "text": "\n".join(line["text"] for line in lines),
                    "method": "ocr",
                    "ocr_lines": lines,
                })
    else:
        with Image.open(path) as image:
            lines = extract_image_text(image, ocr_engine)
        documents.append({
            "page": 1,
            "text": "\n".join(line["text"] for line in lines),
            "method": "ocr",
            "ocr_lines": lines,
        })
    return documents


def _ingest_csv(path: Path) -> list[dict]:
    """Convert CSV headers and rows into chunkable text."""
    with path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        reader = csv.DictReader(csv_file)
        if not reader.fieldnames:
            raise ValueError("CSV file must contain a header row.")

        rows = []
        for row_number, row in enumerate(reader, start=1):
            values = [
                f"{column}: {value.strip()}"
                for column, value in row.items()
                if column and value and value.strip()
            ]
            if values:
                rows.append(f"Row {row_number}\n" + "\n".join(values))

    if not rows:
        raise ValueError("CSV file contains no data rows.")
    return [{"page": 1, "text": "\n\n".join(rows), "method": "csv_text"}]


def _ingest_docx(path: Path) -> list[dict]:
    """Extract paragraphs and table contents from a Word document."""
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX ingestion requires python-docx.") from exc

    document = Document(path)
    sections = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                sections.append(" | ".join(cells))

    if not sections:
        raise ValueError("DOCX file contains no readable text.")
    return [{"page": 1, "text": "\n".join(sections), "method": "docx_text"}]
